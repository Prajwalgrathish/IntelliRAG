import json

import pytest

from src.ingestion import SUPPORTED_EXTENSIONS, get_loader, load_document
from src.ingestion.base import DocumentLoadError
from src.ingestion.service import ingest_directory, ingest_file, list_processed_documents, load_processed_chunks
from src.nlp.chunking import SentenceChunker


def make_pdf(text: str) -> bytes:
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
    ]
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def test_txt_loader_reads_text_and_metadata(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("Hello world. This is a note.", encoding="utf-8")
    doc = load_document(path)
    assert doc.name == "note.txt"
    assert doc.file_type == "txt"
    assert doc.segments[0].text.startswith("Hello world")
    assert doc.metadata["encoding"] == "utf-8-sig"


def test_txt_loader_falls_back_to_cp1252(tmp_path):
    path = tmp_path / "legacy.txt"
    path.write_bytes("caf\xe9 au lait".encode("cp1252"))
    assert "café" in load_document(path).full_text


def test_missing_empty_and_unsupported_files_raise_structured_errors(tmp_path):
    with pytest.raises(DocumentLoadError) as missing:
        load_document(tmp_path / "nope.txt")
    assert missing.value.code == "file_not_found"

    empty = tmp_path / "empty.txt"
    empty.write_bytes(b"")
    with pytest.raises(DocumentLoadError) as empty_err:
        load_document(empty)
    assert empty_err.value.code == "empty_file"

    whitespace = tmp_path / "blank.txt"
    whitespace.write_text("   \n\n  ")
    with pytest.raises(DocumentLoadError) as blank_err:
        load_document(whitespace)
    assert blank_err.value.code == "no_extractable_text"

    odd = tmp_path / "data.xyz"
    odd.write_text("content")
    with pytest.raises(DocumentLoadError) as unsupported:
        get_loader(odd)
    assert unsupported.value.code == "unsupported_type"
    assert {"code", "path", "reason"} <= unsupported.value.to_dict().keys()


def test_pdf_loader_extracts_text_with_page_numbers(tmp_path):
    path = tmp_path / "doc.pdf"
    path.write_bytes(make_pdf("Neural retrieval systems"))
    doc = load_document(path)
    assert doc.file_type == "pdf"
    assert doc.segments[0].page == 1
    assert "Neural retrieval systems" in doc.segments[0].text
    assert doc.metadata["num_pages"] == 1


def test_corrupt_pdf_reports_error(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"this is definitely not a pdf")
    with pytest.raises(DocumentLoadError) as err:
        load_document(path)
    assert err.value.code == "corrupt_pdf"


def test_docx_loader_splits_sections_by_heading_and_reads_tables(tmp_path):
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Transformers power modern NLP.")
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Recall improved after reranking.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "metric"
    table.rows[0].cells[1].text = "value"
    path = tmp_path / "report.docx"
    doc.save(path)

    loaded = load_document(path)
    sections = [s.section for s in loaded.segments]
    assert sections == ["Introduction", "Results"]
    assert "metric | value" in loaded.segments[1].text


def test_corrupt_docx_reports_error(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"not a zip file")
    with pytest.raises(DocumentLoadError) as err:
        load_document(path)
    assert err.value.code == "corrupt_docx"


def test_ingest_file_writes_processed_chunks_and_metadata(tmp_path):
    src = tmp_path / "topic.txt"
    src.write_text(" ".join(f"Sentence number {i} talks about vector search." for i in range(40)))
    processed = tmp_path / "processed"
    result = ingest_file(src, processed, SentenceChunker(50, 10))

    assert result.chunks > 1
    chunks = load_processed_chunks(processed)
    assert len(chunks) == result.chunks
    assert all(c.document == "topic.txt" and c.source_path == str(src) for c in chunks)
    assert len({c.chunk_id for c in chunks}) == len(chunks)
    meta = list_processed_documents(processed)
    assert meta[0]["chunks"] == result.chunks and meta[0]["document"] == "topic.txt"
    assert json.loads((processed / f"{result.doc_id}.meta.json").read_text())["content_sha256"]


def test_reingesting_changed_file_replaces_stale_version(tmp_path):
    src = tmp_path / "doc.txt"
    processed = tmp_path / "processed"
    src.write_text("First version of the document about cats.")
    first = ingest_file(src, processed, SentenceChunker(50, 10))
    src.write_text("Second version of the document about dogs.")
    second = ingest_file(src, processed, SentenceChunker(50, 10))
    assert first.doc_id != second.doc_id
    assert [m["doc_id"] for m in list_processed_documents(processed)] == [second.doc_id]
    assert not (processed / f"{first.doc_id}.jsonl").exists()


def test_ingest_directory_reports_failures_instead_of_ignoring(tmp_path):
    docs = tmp_path / "docs"
    docs.mkdir()
    (docs / "good.txt").write_text("A perfectly fine document about retrieval.")
    (docs / "bad.pdf").write_bytes(b"garbage")
    (docs / "image.png").write_bytes(b"\x89PNG")
    results, errors = ingest_directory(docs, tmp_path / "processed", SentenceChunker(50, 10))
    assert [r.document for r in results] == ["good.txt"]
    assert {e["code"] for e in errors} == {"corrupt_pdf", "unsupported_type"}


def test_supported_extensions_cover_required_types():
    assert {".txt", ".pdf", ".docx"} <= set(SUPPORTED_EXTENSIONS)
